#!/usr/bin/env python3
"""Deep check empty docx file."""
from pathlib import Path
from docx import Document
p = Path("D:/InspurCode/zs-deepseek-harness/non-functional/新建 DOCX 文档.docx")
doc = Document(p)
print("paragraphs:", len(doc.paragraphs))
print("tables:", len(doc.tables))
print("sections:", len(doc.sections))
for i, p in enumerate(doc.paragraphs[:30]):
    print(f"  [{i}] style={p.style.name!r} text={p.text!r}")
for ti, t in enumerate(doc.tables[:5]):
    print(f"table {ti}: rows={len(t.rows)} cols={len(t.columns)}")
    for ri, row in enumerate(t.rows[:5]):
        for ci, cell in enumerate(row.cells[:5]):
            print(f"  [{ri},{ci}] {cell.text!r}")
